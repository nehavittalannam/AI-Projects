from docx import Document
from docx.shared import Pt
import base64
import requests
from openai import AzureOpenAI
import os
from dotenv import load_dotenv

load_dotenv()

# Ensure Images Folder Exists
def ensure_images_folder(folder_name="images"):
    if not os.path.exists(folder_name):
        os.makedirs(folder_name)
 
# Save Image to Disk
def save_image(image_bytes, image_index, folder_name="images"):
    ensure_images_folder(folder_name)
    image_path = os.path.join(folder_name, f"image_{image_index}.jpg")
    with open(image_path, "wb") as f:
        f.write(image_bytes)
    return image_path
 
# Get Image Summary from LLM
def get_image_summary(image_bytes, llm):
    """Send image to LLM for summarization and return the generated text"""
    doc_message = '''Get caption along with a short summary of the image content.
    output format:
         image caption:
         image description
    '''
    
    # Convert image to Base64
    image_base64 = base64.b64encode(image_bytes).decode("utf-8")
    image_url = f"data:image/jpeg;base64,{image_base64}"
    
    # Prepare request payload
    response = llm.chat.completions.create(
        model="your_llm_model",  # Update if needed
        messages=[
            {"role": "system", "content": doc_message},
            {"role": "user", "content": [{"type": "image_url", "image_url": {"url": image_url}}]}
        ]
    )
    
    return response.choices[0].message.content
def get_images(doc):
    """Extract and save images from the document in order and return their paths."""
    image_paths = []
    image_count = 0  # To ensure ordered numbering
 
    for shape in doc.inline_shapes:
        try:
            # Extract the image reference ID
            image_rel_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            
            # Get the image part using the reference ID
            image_part = doc.part.rels[image_rel_id].target_part
            
            # Read image bytes
            image_bytes = image_part.blob
            
            # Save the image
            image_count += 1
            image_path = save_image(image_bytes, image_count)
            image_paths.append(image_path)
        except Exception as e:
            print(f"Error extracting image: {e}")
 
    return image_paths
 
# Replace Images with Summaries in the Correct Order
def replace_images_with_text(doc_path, output_path):
    doc = Document(doc_path)
    image_paths = get_images(doc)  # Extract images in order
    image_index = 0  # Track image index
    
    for para in doc.paragraphs:
        for run in para.runs:
            if "graphicData" in run._element.xml:  # Identify image placeholders
                run._element.clear()  # Remove image
                
                if image_index < len(image_paths):  # Ensure we don't go out of bounds
                    image_summary = get_image_summary(open(image_paths[image_index], 'rb').read(), llm)
                    run.text = f"{image_summary}"
                    image_index += 1
                
                run.font.size = Pt(12)  # Adjust font size if needed
 
    doc.save(output_path)  # Save modified document
 
# Example usage
replace_images_with_text("input_doc", "output_doc")
